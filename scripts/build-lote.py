# -*- coding: utf-8 -*-
"""
Gera a planilha de importacao e o relatorio HTML de um lote de questoes.

Uso: python3 scripts/build-lote.py <lote.json> <saida_sem_extensao> <nivel> <titulo> [pasta_figuras]
"""
import json, html, sys, os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

lote, saida, nivel, titulo = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]
figs = sys.argv[5] if len(sys.argv) > 5 else None
qs = json.load(open(lote))
urls = json.load(open(os.path.join(figs, '_urls.json'))) if figs and os.path.exists(os.path.join(figs, '_urls.json')) else {}

COLS = ['docId','prompt_text','imageUrl','optionA_text','optionA_imageUrl','optionB_text','optionB_imageUrl',
        'optionC_text','optionC_imageUrl','optionD_text','optionD_imageUrl','optionE_text','optionE_imageUrl',
        'correctOptionId','shuffleOptions','explanation','explanationSource','reference','themes',
        'prova_tipo','prova_ano','nivel','Prova','isActive','internalNote']
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'firebase_import'; ws.append(COLS)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2F5496')
ano = int(qs[0]['id'].split('_')[1])
for q in qs:
    img = urls.get(q['img'] + '.png') if q.get('img') else None
    ws.append([q['id'], q['p'], img, q['o'][0], None, q['o'][1], None, q['o'][2], None, q['o'][3], None, None, None,
               q['gab'], 1, q['e'], 'ia', q['r'], q['tema'], 'ME', ano, nivel, None, 1, q.get('nota')])
for k, v in {'A':24,'B':78,'C':30,'D':44,'F':44,'H':44,'J':44,'N':6,'O':6,'P':78,'Q':8,'R':50,'S':34,'Y':52}.items():
    ws.column_dimensions[k].width = v
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(saida + '.xlsx')

L = 'ABCD'
p = ['<meta charset="utf-8"><title>' + html.escape(titulo) + '</title>',
 '<style>body{font-family:-apple-system,Segoe UI,Roboto,sans-serif;max-width:920px;margin:2rem auto;padding:0 1rem;line-height:1.55;color:#1a1a1a}',
 'h1{font-size:1.45rem}h2{font-size:1.05rem;margin-top:2.2rem;border-bottom:2px solid #2F5496;padding-bottom:.3rem}',
 '.q{border:1px solid #ddd;border-radius:8px;padding:1.1rem;margin:1.2rem 0}',
 '.tag{display:inline-block;background:#eef2ff;color:#3730a3;font-size:.72rem;padding:.14rem .5rem;border-radius:4px;margin-right:.35rem}',
 '.op{margin:.28rem 0;padding:.38rem .6rem;border-radius:5px;background:#fafafa}.ok{background:#e7f6ec;font-weight:600}',
 '.ref{font-size:.82rem;color:#555;border-top:1px solid #eee;margin-top:.9rem;padding-top:.5rem}',
 '.ex{background:#fcfcfc;border-left:3px solid #2F5496;padding:.7rem 1rem;margin-top:.7rem;font-size:.94rem}',
 '.nota{background:#fff8e1;border-left:3px solid #f59e0b;padding:.6rem .9rem;margin-top:.6rem;font-size:.85rem}',
 'pre{white-space:pre-wrap;font-family:inherit;margin:.4rem 0}img{max-width:100%;border:1px solid #ddd;border-radius:6px;margin:.5rem 0}</style>',
 f'<h1>{html.escape(titulo)}</h1><p>{len(qs)} questões. Nível {nivel}.</p>']
atual = None
for q in qs:
    prova = q['id'].split('_Q')[0]
    if prova != atual:
        atual = prova; p.append(f'<h2>{prova}</h2>')
    p.append(f'<div class="q"><span class="tag">{html.escape(q["tema"])}</span><span class="tag">{q["id"]}</span>')
    p.append('<pre>' + html.escape(q['p']) + '</pre>')
    if q.get('img') and urls.get(q['img'] + '.png'):
        p.append(f'<img src="{urls[q["img"] + ".png"]}" alt="figura">')
    for j, o in enumerate(q['o']):
        p.append(f'<div class="{"op ok" if L[j] == q["gab"] else "op"}">{L[j]}) {html.escape(o)}</div>')
    p.append(f'<div class="ex">{q["e"]}</div>')
    if q.get('nota'): p.append(f'<div class="nota"><b>Nota interna:</b> {html.escape(q["nota"])}</div>')
    p.append(f'<div class="ref">{html.escape(q["r"])}</div></div>')
open(saida + '.html', 'w', encoding='utf-8').write('\n'.join(p))
print(f'gerados: {saida}.xlsx e {saida}.html ({len(qs)} questoes)')
