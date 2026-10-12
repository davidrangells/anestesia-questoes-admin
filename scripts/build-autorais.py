# -*- coding: utf-8 -*-
"""
Junta lotes de questoes autorais do Quest (JSON) numa planilha de importacao e num relatorio HTML.

Uso: python3 scripts/build-autorais.py <saida_sem_extensao> <lote.json> [lote.json ...]
"""
import json, html, sys, re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

saida, lotes = sys.argv[1], sys.argv[2:]
qs = []
for f in lotes:
    qs += json.load(open(f, encoding='utf-8'))
qs.sort(key=lambda q: int(re.sub(r'\D', '', q['id'])))

COLS = ['docId', 'prompt_text', 'imageUrl', 'optionA_text', 'optionA_imageUrl', 'optionB_text', 'optionB_imageUrl',
        'optionC_text', 'optionC_imageUrl', 'optionD_text', 'optionD_imageUrl', 'optionE_text', 'optionE_imageUrl',
        'correctOptionId', 'shuffleOptions', 'explanation', 'explanationSource', 'reference', 'themes',
        'prova_tipo', 'prova_ano', 'nivel', 'Prova', 'isActive']
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'firebase_import'; ws.append(COLS)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2F5496')
for q in qs:
    o = q['o']
    ws.append([q['id'], q['p'], None, o[0], None, o[1], None, o[2], None, o[3], None, None, None,
               q['gab'], 1, q['e'], 'ia', q['r'], q['tema'], 'QUEST', 2026, q['nivel'], 'Quest', 1])
for k, v in {'A': 16, 'B': 78, 'D': 40, 'F': 40, 'H': 40, 'J': 40, 'N': 6, 'O': 6, 'P': 80, 'R': 46, 'S': 30, 'V': 6, 'W': 18}.items():
    ws.column_dimensions[k].width = v
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(saida + '.xlsx')

L = 'ABCD'
p = ['<meta charset="utf-8"><title>Questões Quest</title>',
     '<style>body{font-family:-apple-system,Segoe UI,Roboto,sans-serif;max-width:940px;margin:2rem auto;padding:0 1rem;line-height:1.55;color:#1a1a1a}',
     'h1{font-size:1.45rem}h2{font-size:1.05rem;margin-top:2.2rem;border-bottom:2px solid #2F5496;padding-bottom:.3rem}',
     '.q{border:1px solid #ddd;border-radius:8px;padding:1.1rem;margin:1.2rem 0}',
     '.tag{display:inline-block;background:#eef2ff;color:#3730a3;font-size:.72rem;padding:.14rem .5rem;border-radius:4px;margin-right:.35rem}',
     '.op{margin:.28rem 0;padding:.38rem .6rem;border-radius:5px;background:#fafafa}.ok{background:#e7f6ec;font-weight:600}',
     '.ref{font-size:.82rem;color:#555;border-top:1px solid #eee;margin-top:.9rem;padding-top:.5rem}',
     '.ex{background:#fcfcfc;border-left:3px solid #2F5496;padding:.7rem 1rem;margin-top:.7rem;font-size:.94rem}</style>',
     f'<h1>Questões Quest</h1><p>{len(qs)} questões.</p>']
atual = None
for q in qs:
    if q['tema'] != atual:
        atual = q['tema']; p.append(f'<h2>{html.escape(atual)}</h2>')
    p.append(f'<div class="q"><span class="tag">{q["id"]}</span><span class="tag">{q["nivel"]}</span>')
    p.append('<p>' + html.escape(q['p']) + '</p>')
    for j, o in enumerate(q['o']):
        p.append(f'<div class="{"op ok" if L[j] == q["gab"] else "op"}">{L[j]}) {html.escape(o)}</div>')
    p.append(f'<div class="ex">{q["e"]}</div>')
    p.append(f'<div class="ref">{html.escape(q["r"])}</div></div>')
open(saida + '.html', 'w', encoding='utf-8').write('\n'.join(p))
print(f'gerados: {saida}.xlsx e {saida}.html ({len(qs)} questoes)')
